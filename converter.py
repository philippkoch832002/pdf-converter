import os
from docx import Document
from fpdf import FPDF

def docx_to_pdf(input_path, output_path):
    doc = Document(input_path)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.add_page();
    pdf.set_font("Arial", size=12)

    for para in doc.paragraphs:
        pdf.cell(200, 10, txt=para.text, ln=True)

    pdf.output(output_path)
    return output_path    